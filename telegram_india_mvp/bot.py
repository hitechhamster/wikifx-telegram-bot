import asyncio
import csv
import html
import logging
import os
import re
import sqlite3
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from urllib.parse import urlparse
from uuid import uuid4

from openpyxl import load_workbook
from telegram import (
    BotCommand,
    BotCommandScopeAllGroupChats,
    BotCommandScopeAllPrivateChats,
    BotCommandScopeChat,
    BotCommandScopeChatMember,
    BotCommandScopeDefault,
    ChatPermissions,
    ForceReply,
    InlineKeyboardButton,
    InlineKeyboardMarkup,
    InlineQueryResultArticle,
    InputTextMessageContent,
    Update,
)
from telegram.constants import ChatAction, ParseMode
from telegram.error import TelegramError
from telegram.ext import (
    Application,
    CallbackQueryHandler,
    CommandHandler,
    ContextTypes,
    InlineQueryHandler,
    MessageHandler,
    filters,
)

from ai_agent import AIServiceError, ask_deepseek, is_ai_enabled


PROJECT_DIR = Path(__file__).resolve().parent
DEFAULT_XLSX = PROJECT_DIR.parent / "交易商名单.xlsx"
DEFAULT_CURATED_XLSX = (
    PROJECT_DIR.parent
    / "outputs"
    / "telegram_india_mvp_20260819"
    / "Telegram_India_Broker_MVP.xlsx"
)
XLSX_PATH = Path(os.getenv("BROKER_XLSX_PATH", str(DEFAULT_XLSX)))
CURATED_XLSX_PATH = Path(
    os.getenv("CURATED_BROKER_XLSX_PATH", str(DEFAULT_CURATED_XLSX))
)
DB_PATH = Path(os.getenv("BOT_DB_PATH", str(PROJECT_DIR / "telegram_mvp.db")))
RETENTION_DAYS = int(os.getenv("RETENTION_DAYS", "30"))
MARKET = os.getenv("BOT_MARKET", "India")
SOURCE_TAG = os.getenv("SOURCE_TAG", "india_own_group_mvp")
WIKIFX_APP_ONELINK = os.getenv(
    "WIKIFX_APP_ONELINK",
    "https://fxeye.onelink.me/Vm4A/intgbot",
)
CHECK_PROMPT = (
    "Reply with a broker name, official website or WikiFX ID.\n"
    "Example: XM or xm.com"
)
PUBLIC_COMMANDS = (
    BotCommand("start", "How to use the bot"),
    BotCommand("check", "Find a broker"),
)
MAX_RESULTS = 5
RECALL_AFTER_DAYS = int(os.getenv("RECALL_AFTER_DAYS", "3"))
RECALL_REPEAT_DAYS = int(os.getenv("RECALL_REPEAT_DAYS", "7"))
RECALL_CHECK_SECONDS = int(os.getenv("RECALL_CHECK_SECONDS", "21600"))

logging.basicConfig(
    format="%(asctime)s %(levelname)s %(name)s %(message)s",
    level=logging.INFO,
)
# httpx logs include the full Telegram Bot API URL, which contains the bot token.
# Keep request-level logging disabled so terminal output and screenshots do not
# accidentally expose credentials.
logging.getLogger("httpx").setLevel(logging.WARNING)
LOGGER = logging.getLogger("wikifx-telegram-mvp")
BROKER_CACHE = None
BROKER_INDEXES = None


def admin_ids():
    raw = os.getenv("TELEGRAM_ADMIN_IDS", "")
    return {int(value.strip()) for value in raw.split(",") if value.strip().isdigit()}


def is_admin(update: Update):
    user = update.effective_user
    return bool(user and user.id in admin_ids())


def connect_db():
    connection = sqlite3.connect(DB_PATH)
    connection.row_factory = sqlite3.Row
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            event_time TEXT NOT NULL,
            event_type TEXT NOT NULL,
            telegram_user_id INTEGER,
            username TEXT,
            chat_id INTEGER,
            chat_type TEXT,
            market TEXT,
            query_text TEXT,
            matched_broker_id TEXT,
            match_type TEXT,
            result_count INTEGER,
            source_tag TEXT
        )
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS broker_follows (
            telegram_user_id INTEGER NOT NULL,
            private_chat_id INTEGER NOT NULL,
            broker_id TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL,
            active INTEGER NOT NULL DEFAULT 1,
            last_news_at TEXT,
            last_recall_at TEXT,
            PRIMARY KEY (telegram_user_id, broker_id)
        )
        """
    )
    connection.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_broker_follows_active_broker
        ON broker_follows (active, broker_id)
        """
    )
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS ai_messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            event_time TEXT NOT NULL,
            telegram_user_id INTEGER NOT NULL,
            chat_id INTEGER NOT NULL,
            role TEXT NOT NULL,
            content TEXT NOT NULL
        )
        """
    )
    connection.execute(
        """
        CREATE INDEX IF NOT EXISTS idx_ai_messages_conversation
        ON ai_messages (telegram_user_id, chat_id, id)
        """
    )
    connection.commit()
    return connection


DB = connect_db()


def cleanup_old_events():
    cutoff = datetime.now(timezone.utc) - timedelta(days=RETENTION_DAYS)
    DB.execute("DELETE FROM events WHERE event_time < ?", (cutoff.isoformat(),))
    DB.execute("DELETE FROM ai_messages WHERE event_time < ?", (cutoff.isoformat(),))
    DB.commit()


def log_event(
    update: Update,
    event_type: str,
    query_text: str = "",
    broker_id: str = "",
    match_type: str = "",
    result_count: int = 0,
):
    user = update.effective_user
    chat = update.effective_chat
    DB.execute(
        """
        INSERT INTO events (
            event_time, event_type, telegram_user_id, username,
            chat_id, chat_type, market, query_text,
            matched_broker_id, match_type, result_count, source_tag
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            datetime.now(timezone.utc).isoformat(),
            event_type,
            user.id if user else None,
            user.username if user else None,
            chat.id if chat else None,
            chat.type if chat else None,
            MARKET,
            query_text,
            broker_id,
            match_type,
            result_count,
            SOURCE_TAG,
        ),
    )
    DB.commit()


def follow_status(user_id: int, broker_id: str) -> bool:
    row = DB.execute(
        """
        SELECT active FROM broker_follows
        WHERE telegram_user_id = ? AND broker_id = ?
        """,
        (user_id, str(broker_id)),
    ).fetchone()
    return bool(row and row["active"])


def save_follow(user_id: int, private_chat_id: int, broker_id: str, active: bool):
    now = datetime.now(timezone.utc).isoformat()
    DB.execute(
        """
        INSERT INTO broker_follows (
            telegram_user_id, private_chat_id, broker_id,
            created_at, updated_at, active
        ) VALUES (?, ?, ?, ?, ?, ?)
        ON CONFLICT(telegram_user_id, broker_id) DO UPDATE SET
            private_chat_id = excluded.private_chat_id,
            updated_at = excluded.updated_at,
            active = excluded.active
        """,
        (user_id, private_chat_id, str(broker_id), now, now, int(active)),
    )
    DB.commit()


def active_follows(user_id: int):
    return DB.execute(
        """
        SELECT * FROM broker_follows
        WHERE telegram_user_id = ? AND active = 1
        ORDER BY updated_at DESC
        """,
        (user_id,),
    ).fetchall()


def welcome_text():
    return (
        "<b>WikiFX Broker Check India</b> 🔎\n\n"
        "Check a broker’s registered region, website, regulation summary and WikiFX profile before you trade.\n\n"
        "<b>How it works</b>\n"
        "1. Tap <b>Check a broker</b>.\n"
        "2. Send a broker name, official website or WikiFX ID.\n"
        "3. Tap <b>Follow updates</b> on the result card to receive future news and reminders.\n\n"
        "<i>Powered by WikiFX broker data and risk intelligence.</i>"
    )


def welcome_buttons(is_private: bool, bot_username: str):
    follows_button = (
        InlineKeyboardButton("⭐ My followed brokers", callback_data="nav:follows")
        if is_private
        else InlineKeyboardButton(
            "⭐ My followed brokers",
            url=f"https://t.me/{bot_username}?start=my_follows",
        )
    )
    return InlineKeyboardMarkup(
        [
            [InlineKeyboardButton("🔎 Check a broker", callback_data="nav:check")],
            [follows_button],
            [InlineKeyboardButton("📲 Open WikiFX App", url=WIKIFX_APP_ONELINK)],
        ]
    )


def follows_page(user_id: int):
    rows = []
    names = []
    for follow in active_follows(user_id)[:10]:
        records, _ = find_brokers(follow["broker_id"])
        if not records:
            continue
        broker = records[0]
        broker_name = text_value(broker.get("broker_name")) or "Broker"
        names.append(f"• {html.escape(broker_name)} · ID {html.escape(str(broker['broker_id']))}")
        rows.append(
            [
                InlineKeyboardButton(
                    f"🔎 {broker_name[:24]}",
                    callback_data=f"followview:{broker['broker_id']}",
                )
            ]
        )

    if names:
        text = "<b>My followed brokers</b> ⭐\n\n" + "\n".join(names)
    else:
        text = (
            "<b>My followed brokers</b> ⭐\n\n"
            "You are not following any brokers yet. Check a broker and tap Follow updates."
        )
    rows.extend(
        [
            [InlineKeyboardButton("🔎 Check another broker", callback_data="nav:check")],
            [InlineKeyboardButton("🏠 Home", callback_data="nav:home")],
        ]
    )
    return text, InlineKeyboardMarkup(rows)


def split_terms(value):
    if value is None:
        return []
    return [str(item).strip().casefold() for item in str(value).split("|") if str(item).strip()]


def text_value(value):
    if value is None:
        return ""
    return str(value).strip()


def broker_id_value(value):
    raw = text_value(value)
    if raw.endswith(".0") and raw[:-2].isdigit():
        raw = raw[:-2]
    return raw.zfill(10) if raw.isdigit() and len(raw) <= 10 else raw


def date_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return text_value(value)


def numeric_score(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return -1.0


def unique_terms(*groups):
    seen = set()
    output = []
    for group in groups:
        for item in group:
            term = text_value(item).casefold()
            if term and term not in seen:
                seen.add(term)
                output.append(term)
    return output


def license_lookup_terms(licenses):
    """Return whole licence components such as 335692, 372/18 or SD018."""
    return {
        term
        for licence in licenses
        for term in re.findall(r"[a-z0-9]+(?:[./-][a-z0-9]+)*", licence)
    }


def normalise_query(query):
    value = re.sub(r"\s+", " ", query.strip().casefold())
    if "://" in value:
        parsed = urlparse(value)
        dealer_id = re.search(r"/dealer/(\d+)", parsed.path)
        if dealer_id:
            return broker_id_value(dealer_id.group(1))
        value = parsed.netloc or parsed.path
    value = value.removeprefix("www.").split("/", 1)[0].strip("/ ")
    return value


def load_curated_brokers(path):
    if not path.exists():
        return []
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "India Brokers" not in workbook.sheetnames:
        workbook.close()
        return []
    sheet = workbook["India Brokers"]
    rows = sheet.iter_rows(min_row=4, values_only=True)
    headers = [str(value).strip() for value in next(rows)]
    records = []
    for values in rows:
        record = dict(zip(headers, values))
        if not record.get("broker_id"):
            continue
        if str(record.get("active", True)).strip().casefold() in {"false", "0", "no"}:
            continue
        record["broker_id"] = broker_id_value(record["broker_id"])
        record["broker_name"] = text_value(record.get("broker_name"))
        record["full_name"] = record["broker_name"]
        record["aliases_list"] = split_terms(record.get("aliases"))
        record["domains_list"] = split_terms(record.get("official_domains"))
        record["licenses_list"] = split_terms(record.get("license_numbers"))
        record["license_lookup_terms"] = license_lookup_terms(record["licenses_list"])
        record["country"] = text_value(record.get("market"))
        record["regulatory_summary"] = text_value(record.get("regulation_status"))
        record["regulatory_warning"] = text_value(record.get("risk_summary"))
        record["risk_grade"] = ""
        record["score"] = ""
        record["source_checked_at"] = date_value(record.get("source_checked_at"))
        record["has_curated_identity"] = True
        records.append(record)
    workbook.close()
    return records


def load_full_brokers(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "交易商名单" not in workbook.sheetnames:
        workbook.close()
        return []
    sheet = workbook["交易商名单"]
    rows = sheet.iter_rows(min_row=1, values_only=True)
    headers = [text_value(value) for value in next(rows)]
    records = []
    for values in rows:
        source = dict(zip(headers, values))
        broker_id = broker_id_value(source.get("WikiFX ID"))
        full_name = text_value(source.get("全称"))
        broker_name = text_value(source.get("名称")) or full_name
        if not broker_id:
            continue
        if not broker_name:
            broker_name = f"WikiFX Broker {broker_id}"
        country = text_value(source.get("国家/地区"))
        regulatory_summary = text_value(source.get("监管概要"))
        regulatory_warning = text_value(source.get("监管警告"))
        licence_count = text_value(source.get("牌照数"))
        regulation_status = regulatory_summary
        if not regulation_status and licence_count:
            regulation_status = f"{licence_count} regulatory record(s) shown on WikiFX"
        wikifx_url = text_value(source.get("WikiFX 页面"))
        if not wikifx_url:
            wikifx_url = f"https://www.wikifx.com/en/dealer/{broker_id}.html"
        aliases = unique_terms([broker_name])
        records.append(
            {
                "broker_id": broker_id,
                "broker_name": broker_name,
                "full_name": full_name,
                "aliases": "|".join(aliases),
                "aliases_list": aliases,
                "official_domains": "",
                "domains_list": [],
                "license_numbers": "",
                "licenses_list": [],
                "license_lookup_terms": set(),
                "regulation_status": regulation_status,
                "regulatory_summary": regulatory_summary,
                "regulatory_warning": regulatory_warning,
                "risk_summary": regulatory_warning,
                "complaint_count": "",
                "wikifx_url": wikifx_url,
                "app_download_url": WIKIFX_APP_ONELINK,
                "market": country,
                "country": country,
                "language": "English",
                "source_checked_at": date_value(source.get("最后抓取")),
                "active": True,
                "has_license": text_value(source.get("有牌照")),
                "license_count": licence_count,
                "score": source.get("评分"),
                "risk_grade": text_value(source.get("风险档")),
                "source": text_value(source.get("来源")),
            }
        )
    workbook.close()
    return records


def merge_curated_data(records, curated_records):
    by_id = {record["broker_id"]: record for record in records}
    for curated in curated_records:
        broker = by_id.get(curated["broker_id"])
        if broker is None:
            records.append(curated)
            by_id[curated["broker_id"]] = curated
            continue
        broker["aliases_list"] = unique_terms(
            broker.get("aliases_list", []), curated.get("aliases_list", [])
        )
        broker["domains_list"] = unique_terms(
            broker.get("domains_list", []), curated.get("domains_list", [])
        )
        broker["licenses_list"] = unique_terms(
            broker.get("licenses_list", []), curated.get("licenses_list", [])
        )
        broker["license_lookup_terms"] = license_lookup_terms(broker["licenses_list"])
        broker["official_domains"] = "|".join(broker["domains_list"])
        broker["license_numbers"] = "|".join(broker["licenses_list"])
        broker["has_curated_identity"] = True
    return records


def build_broker_indexes(records):
    indexes = {
        "broker_id": defaultdict(list),
        "name": defaultdict(list),
        "full_name": defaultdict(list),
        "domain": defaultdict(list),
        "license": defaultdict(list),
    }
    for broker in records:
        indexes["broker_id"][broker["broker_id"].casefold()].append(broker)
        name_terms = unique_terms([broker.get("broker_name", "")], broker.get("aliases_list", []))
        full_name_terms = unique_terms([broker.get("full_name", "")])
        broker["search_name_terms"] = unique_terms(name_terms, full_name_terms)
        for term in name_terms:
            indexes["name"][term].append(broker)
        for term in full_name_terms:
            indexes["full_name"][term].append(broker)
        for term in broker.get("domains_list", []):
            indexes["domain"][term].append(broker)
        for term in unique_terms(
            broker.get("licenses_list", []), broker.get("license_lookup_terms", set())
        ):
            indexes["license"][term].append(broker)
    return indexes


def load_brokers():
    global BROKER_CACHE, BROKER_INDEXES
    if BROKER_CACHE is not None:
        return BROKER_CACHE
    if not XLSX_PATH.exists():
        raise FileNotFoundError(f"Broker workbook not found: {XLSX_PATH}")

    workbook = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    sheet_names = set(workbook.sheetnames)
    workbook.close()
    if "交易商名单" in sheet_names:
        records = load_full_brokers(XLSX_PATH)
        records = merge_curated_data(records, load_curated_brokers(CURATED_XLSX_PATH))
    elif "India Brokers" in sheet_names:
        records = load_curated_brokers(XLSX_PATH)
    else:
        raise ValueError("Broker workbook must contain '交易商名单' or 'India Brokers'.")

    BROKER_CACHE = records
    BROKER_INDEXES = build_broker_indexes(records)
    return BROKER_CACHE


def sorted_brokers(records):
    unique = {record["broker_id"]: record for record in records}
    return sorted(
        unique.values(),
        key=lambda record: (
            not bool(record.get("has_curated_identity")),
            not bool(record.get("domains_list")),
            not bool(text_value(record.get("regulatory_summary"))),
            text_value(record.get("broker_name")).casefold(),
            text_value(record.get("country")).casefold(),
            record["broker_id"],
        ),
    )


def find_brokers(query):
    keyword = normalise_query(query)
    if not keyword:
        return [], "empty"
    load_brokers()
    for match_type in ("broker_id", "name", "domain", "license", "full_name"):
        exact = BROKER_INDEXES[match_type].get(keyword, [])
        if exact:
            return sorted_brokers(exact)[:MAX_RESULTS], match_type

    if len(keyword) < 2:
        return [], "none"
    partial = []
    for broker in BROKER_CACHE:
        searchable = broker.get("search_name_terms", []) + broker.get("domains_list", [])
        if any(keyword in term for term in searchable):
            partial.append(broker)
    partial.sort(
        key=lambda broker: (
            not any(term.startswith(keyword) for term in broker.get("search_name_terms", [])),
            not bool(broker.get("has_curated_identity")),
            not bool(broker.get("domains_list")),
            not bool(text_value(broker.get("regulatory_summary"))),
            text_value(broker.get("broker_name")).casefold(),
            broker["broker_id"],
        )
    )
    seen = set()
    ordered_partial = []
    for broker in partial:
        if broker["broker_id"] not in seen:
            seen.add(broker["broker_id"])
            ordered_partial.append(broker)
    return ordered_partial[:MAX_RESULTS], "partial" if ordered_partial else "none"


def broker_card(broker):
    def escaped(value, fallback="Not available"):
        return html.escape(text_value(value) or fallback)

    def shortened(value, limit=170):
        text = text_value(value)
        return text if len(text) <= limit else f"{text[: limit - 1].rstrip()}…"

    name = escaped(broker.get("broker_name"))
    country = escaped(broker.get("country"), "Not listed")
    domain = escaped(
        next(iter(broker.get("domains_list", [])), ""),
        "Not listed in the current data",
    )
    regulation = escaped(
        shortened(broker.get("regulatory_summary") or broker.get("license_numbers")),
        "Open the full profile for licence details",
    )
    updated = escaped(broker.get("source_checked_at"), "Not listed")
    broker_id = escaped(broker.get("broker_id"))

    lines = [
        f"<b>🔎 {name}</b>",
        "",
        f"🌍 <b>Country/region:</b> {country}",
        f"🌐 <b>Official website:</b> {domain}",
        f"🆔 <b>WikiFX ID:</b> {broker_id}",
        f"🛡 <b>Regulation:</b> {regulation}",
        f"🕒 <b>Updated:</b> {updated}",
        "",
        "<b>Confirm before following:</b> compare the website domain and regulated entity "
        "with your account-opening email or client agreement.",
        "",
        "Tap below for the full WikiFX profile and latest evidence.",
        "<i>Powered by WikiFX broker data and risk intelligence.</i>",
    ]
    return "\n".join(lines)


def multiple_match_text(query, records):
    lines = [
        f"<b>More than one WikiFX record matches “{html.escape(text_value(query))}”.</b>",
        "",
        "Choose the record whose website matches the domain shown in your broker app, "
        "account-opening email or client agreement.",
        "",
    ]
    for index, record in enumerate(records, start=1):
        name = html.escape(text_value(record.get("broker_name")) or "Broker")
        country = html.escape(text_value(record.get("country")) or "Region not listed")
        domain = next(iter(record.get("domains_list", [])), "")
        regulation = shortened_identity_text(
            record.get("regulatory_summary") or record.get("license_numbers"), 85
        )
        lines.append(f"<b>{index}. {name} · {country}</b>")
        if domain:
            lines.append(
                f"   ✅ Website: <code>{html.escape(domain)}</code> — choose this only if the domain matches"
            )
        else:
            lines.append("   Website: Not listed in the current data")
        lines.append(
            f"   Regulation: {html.escape(regulation or 'Open the profile for licence details')}"
        )
        lines.append(f"   WikiFX ID: <code>{html.escape(str(record['broker_id']))}</code>")
        lines.append("")
    lines.extend(
        [
            "<b>Still unsure?</b> Send the exact official website or 10-digit WikiFX ID. "
            "Use the Profile button to compare the regulated entity and licence details. "
            "Do not choose by brand name or logo alone.",
        ]
    )
    return "\n".join(lines)


def shortened_identity_text(value, limit=85):
    text = text_value(value)
    return text if len(text) <= limit else f"{text[: limit - 1].rstrip()}…"


def broker_buttons(broker, user_id=None, is_private=False, bot_username=""):
    broker_id = str(broker["broker_id"])
    broker_name = text_value(broker.get("broker_name"))[:22] or "broker"
    rows = []

    if is_private and user_id:
        is_following = follow_status(user_id, broker_id)
        rows.append(
            [
                InlineKeyboardButton(
                    (
                        f"✅ Following {broker_name} · tap to unfollow"
                        if is_following
                        else "🔔 Follow for news & risk updates"
                    ),
                    callback_data=(
                        f"unfollow:{broker_id}" if is_following else f"follow:{broker_id}"
                    ),
                )
            ]
        )
        rows.append(
            [InlineKeyboardButton("⭐ My followed brokers", callback_data="nav:follows")]
        )
    elif bot_username:
        rows.append(
            [
                InlineKeyboardButton(
                    "🔔 Follow for news & risk updates",
                    url=f"https://t.me/{bot_username}?start=follow_{broker_id}",
                )
            ]
        )

    rows.extend(
        [
            [
                InlineKeyboardButton("📲 Open App", url=WIKIFX_APP_ONELINK),
                InlineKeyboardButton("🌐 Full profile", url=broker["wikifx_url"]),
            ],
        ]
    )
    return InlineKeyboardMarkup(rows)


def candidate_buttons(records):
    rows = []
    for index, record in enumerate(records, start=1):
        domain = next(iter(record.get("domains_list", [])), "")
        identity = domain or f"ID …{record['broker_id'][-4:]}"
        country = text_value(record.get("country"))[:12] or "Unknown"
        rows.append(
            [
                InlineKeyboardButton(
                    f"{index}. {identity[:24]} · {country}",
                    callback_data=f"broker:{record['broker_id']}",
                ),
                InlineKeyboardButton("Profile ↗", url=record["wikifx_url"]),
            ]
        )
    rows.append(
        [InlineKeyboardButton("🔁 Enter website or WikiFX ID", callback_data="nav:check")]
    )
    return InlineKeyboardMarkup(rows)



def ai_history(user_id: int, chat_id: int):
    rows = DB.execute(
        """
        SELECT role, content FROM ai_messages
        WHERE telegram_user_id = ? AND chat_id = ?
        ORDER BY id DESC LIMIT 8
        """,
        (user_id, chat_id),
    ).fetchall()
    return [
        {"role": row["role"], "content": row["content"]}
        for row in reversed(rows)
        if row["role"] in {"user", "assistant"}
    ]


def save_ai_message(user_id: int, chat_id: int, role: str, content: str):
    DB.execute(
        """
        INSERT INTO ai_messages (
            event_time, telegram_user_id, chat_id, role, content
        ) VALUES (?, ?, ?, ?, ?)
        """,
        (
            datetime.now(timezone.utc).isoformat(),
            user_id,
            chat_id,
            role,
            content[:4000],
        ),
    )
    DB.commit()


def broker_tool_record(broker):
    return {
        "broker_id": str(broker.get("broker_id") or ""),
        "name": text_value(broker.get("broker_name")),
        "company_name": text_value(broker.get("full_name")),
        "country_or_region": text_value(broker.get("country")),
        "official_websites": list(broker.get("domains_list") or []),
        "licence_numbers": list(broker.get("licenses_list") or []),
        "regulation_summary": text_value(
            broker.get("regulatory_summary") or broker.get("regulation_status")
        ),
        "risk_warning": text_value(
            broker.get("regulatory_warning") or broker.get("risk_summary")
        ),
        "wikifx_score": text_value(broker.get("score")),
        "risk_grade": text_value(broker.get("risk_grade")),
        "data_updated_at": text_value(broker.get("source_checked_at")),
        "wikifx_profile_url": text_value(broker.get("wikifx_url")),
        "app_url": text_value(broker.get("app_download_url")) or WIKIFX_APP_ONELINK,
    }


async def execute_ai_tool(name: str, arguments: dict, update: Update):
    if name == "search_brokers":
        query = text_value(arguments.get("query"))[:200]
        records, match_type = find_brokers(query)
        return {
            "ok": True,
            "query": query,
            "match_type": match_type,
            "result_count": len(records),
            "results": [broker_tool_record(record) for record in records],
        }

    if name == "list_my_follows":
        if not update.effective_chat or update.effective_chat.type != "private":
            return {
                "ok": False,
                "error": "Follow management is available only in the bot's private chat.",
            }
        records = []
        for follow in active_follows(update.effective_user.id)[:10]:
            matches, _ = find_brokers(follow["broker_id"])
            if matches:
                records.append(broker_tool_record(matches[0]))
        return {"ok": True, "result_count": len(records), "results": records}

    if name == "set_broker_follow":
        if not update.effective_chat or update.effective_chat.type != "private":
            return {
                "ok": False,
                "error": "Follow management is available only in the bot's private chat.",
            }
        requested_id = broker_id_value(arguments.get("broker_id"))
        action = text_value(arguments.get("action")).casefold()
        records, _ = find_brokers(requested_id)
        broker = next(
            (record for record in records if record["broker_id"] == requested_id),
            None,
        )
        if not broker:
            return {"ok": False, "error": "No exact broker was found for that WikiFX ID."}
        if action not in {"follow", "unfollow"}:
            return {"ok": False, "error": "Action must be follow or unfollow."}

        enabled = action == "follow"
        save_follow(
            update.effective_user.id,
            update.effective_chat.id,
            broker["broker_id"],
            enabled,
        )
        log_event(
            update,
            "follow_enabled" if enabled else "follow_disabled",
            broker_id=broker["broker_id"],
        )
        return {
            "ok": True,
            "action": action,
            "broker": broker_tool_record(broker),
        }

    return {"ok": False, "error": f"Unknown tool: {name}"}


def natural_language_prompt(update: Update, context: ContextTypes.DEFAULT_TYPE):
    message = update.effective_message
    if not message or not message.text or not update.effective_chat:
        return ""

    text = message.text.strip()
    if update.effective_chat.type == "private":
        return text

    username = context.bot.username or ""
    mention_pattern = rf"@{re.escape(username)}\b" if username else ""
    mentioned = bool(mention_pattern and re.search(mention_pattern, text, re.IGNORECASE))
    replied_to_bot = bool(
        message.reply_to_message
        and message.reply_to_message.from_user
        and message.reply_to_message.from_user.id == context.bot.id
    )
    if not mentioned and not replied_to_bot:
        return ""
    if mentioned:
        text = re.sub(mention_pattern, "", text, flags=re.IGNORECASE)
    return text.strip(" ,，:：")


def append_app_download_cta(prompt: str, answer: str):
    if WIKIFX_APP_ONELINK in answer:
        return answer
    if re.search(r"[\u4e00-\u9fff]", prompt):
        cta = (
            "📲 下载 WikiFX App，查看完整交易商资料并持续获取风险动态："
            f"{WIKIFX_APP_ONELINK}"
        )
    else:
        cta = (
            "📲 Download the WikiFX App for full broker details and ongoing risk updates: "
            f"{WIKIFX_APP_ONELINK}"
        )
    return f"{answer.rstrip()}\n\n{cta}"


async def ai_text_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    prompt = natural_language_prompt(update, context)
    if not prompt:
        return

    if not is_ai_enabled():
        await update.effective_message.reply_text(
            "AI service is not configured yet. Use /check <broker name> for the current lookup."
        )
        return

    user = update.effective_user
    chat = update.effective_chat
    if not user or not chat:
        return

    try:
        await context.bot.send_chat_action(chat_id=chat.id, action=ChatAction.TYPING)
    except TelegramError:
        pass

    previous_messages = ai_history(user.id, chat.id)
    log_event(update, "ai_query", query_text=prompt[:1000])

    try:
        answer = await ask_deepseek(
            prompt,
            previous_messages,
            lambda name, args: execute_ai_tool(name, args, update),
        )
    except AIServiceError:
        LOGGER.exception("DeepSeek assistant request failed")
        await update.effective_message.reply_text(
            "The AI assistant is temporarily unavailable. You can still use /check <broker name>."
        )
        return

    answer = answer.replace("**", "").replace("```", "").strip()
    answer = append_app_download_cta(prompt, answer)
    save_ai_message(user.id, chat.id, "user", prompt)
    save_ai_message(user.id, chat.id, "assistant", answer)
    log_event(update, "ai_response")
    await update.effective_message.reply_text(
        answer[:4000],
        disable_web_page_preview=True,
    )


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    log_event(update, "bot_start")

    if context.args and context.args[0] == "my_follows":
        if not update.effective_chat or update.effective_chat.type != "private":
            await update.message.reply_text("Open a private chat with the bot to manage follows.")
            return
        text, markup = follows_page(update.effective_user.id)
        await update.message.reply_text(text, parse_mode=ParseMode.HTML, reply_markup=markup)
        return

    if context.args and context.args[0].startswith("follow_"):
        broker_id = context.args[0].removeprefix("follow_").strip()
        records, _ = find_brokers(broker_id)
        if not records:
            await update.message.reply_text(
                "This broker is no longer active in the current WikiFX workbook."
            )
            return
        if not update.effective_chat or update.effective_chat.type != "private":
            await update.message.reply_text("Open a private chat with the bot to follow updates.")
            return

        broker = records[0]
        save_follow(
            update.effective_user.id,
            update.effective_chat.id,
            broker["broker_id"],
            True,
        )
        log_event(update, "follow_enabled", broker_id=broker["broker_id"])
        await update.message.reply_text(
            f"<b>✅ Following {html.escape(text_value(broker.get('broker_name')))}</b>\n\n"
            "You can now receive new broker news and risk-update reminders in this private chat. "
            "Use the buttons below to check the broker or manage your followed brokers.",
            parse_mode=ParseMode.HTML,
            reply_markup=broker_buttons(
                broker,
                user_id=update.effective_user.id,
                is_private=True,
                bot_username=context.bot.username or "",
            ),
        )
        return

    await update.message.reply_text(
        welcome_text(),
        parse_mode=ParseMode.HTML,
        reply_markup=welcome_buttons(
            bool(update.effective_chat and update.effective_chat.type == "private"),
            context.bot.username or "",
        ),
    )


async def privacy(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update):
        return
    await update.message.reply_text(
        "Privacy notice — internal MVP\n\n"
        "To operate and evaluate this service, we store your Telegram user ID, "
        "username, chat ID, query text, matched broker and timestamps for up to "
        f"{RETENTION_DAYS} days. Do not submit passwords, account numbers, deposits "
        "or identity documents. Public launch requires an approved privacy-policy URL."
    )


async def send_broker_result(update: Update, broker, query_text, match_type):
    log_event(update, "match_success", query_text, broker["broker_id"], match_type, 1)
    is_private = bool(update.effective_chat and update.effective_chat.type == "private")
    user_id = update.effective_user.id if update.effective_user else None
    await update.effective_message.reply_text(
        broker_card(broker),
        parse_mode=ParseMode.HTML,
        reply_markup=broker_buttons(
            broker,
            user_id=user_id,
            is_private=is_private,
            bot_username=update.get_bot().username or "",
        ),
        disable_web_page_preview=True,
    )
    log_event(update, "result_sent", query_text, broker["broker_id"], match_type, 1)


async def process_broker_query(update: Update, query_text: str):
    log_event(update, "broker_query", query_text)
    records, match_type = find_brokers(query_text)

    if not records:
        log_event(update, "match_none", query_text, match_type="none")
        await update.effective_message.reply_text(
            "No match found. Try the exact broker name, official website or 10-digit WikiFX ID."
        )
        return

    if len(records) > 1:
        log_event(update, "match_multiple", query_text, match_type="multiple", result_count=len(records))
        await update.effective_message.reply_text(
            multiple_match_text(query_text, records),
            parse_mode=ParseMode.HTML,
            reply_markup=candidate_buttons(records),
        )
        return

    await send_broker_result(update, records[0], query_text, match_type)


async def check(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query_text = " ".join(context.args).strip()
    if not query_text:
        await update.message.reply_text(
            CHECK_PROMPT,
            reply_markup=ForceReply(
                selective=True,
                input_field_placeholder="Broker name, website or WikiFX ID",
            ),
        )
        return

    await process_broker_query(update, query_text)


async def text_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    message = update.effective_message
    replied_to = message.reply_to_message if message else None
    is_check_reply = bool(
        replied_to
        and replied_to.from_user
        and replied_to.from_user.id == context.bot.id
        and replied_to.text == CHECK_PROMPT
    )
    if is_check_reply:
        query_text = message.text.strip()
        if query_text:
            await process_broker_query(update, query_text)
        return

    await ai_text_message(update, context)


async def candidate_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    callback = update.callback_query
    await callback.answer()
    broker_id = callback.data.split(":", 1)[1]
    records, match_type = find_brokers(broker_id)
    if not records:
        await callback.edit_message_text("This broker is no longer active in the workbook.")
        return
    await callback.edit_message_text("Opening the selected WikiFX record…")
    await send_broker_result(update, records[0], broker_id, match_type)


async def follow_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    callback = update.callback_query
    if not update.effective_chat or update.effective_chat.type != "private":
        await callback.answer(
            "Open a private chat with the bot before following broker updates.",
            show_alert=True,
        )
        return

    action, broker_id = callback.data.split(":", 1)
    records, _ = find_brokers(broker_id)
    if not records:
        await callback.answer("This broker is no longer active.", show_alert=True)
        return

    broker = records[0]
    enabled = action == "follow"
    save_follow(
        update.effective_user.id,
        update.effective_chat.id,
        broker["broker_id"],
        enabled,
    )
    log_event(
        update,
        "follow_enabled" if enabled else "follow_disabled",
        broker_id=broker["broker_id"],
    )
    await callback.answer(
        (
            f"Following {text_value(broker.get('broker_name'))}."
            if enabled
            else f"Updates stopped for {text_value(broker.get('broker_name'))}."
        )
    )
    await callback.edit_message_reply_markup(
        reply_markup=broker_buttons(
            broker,
            user_id=update.effective_user.id,
            is_private=True,
            bot_username=context.bot.username or "",
        )
    )


async def navigation_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    callback = update.callback_query
    action = callback.data.split(":", 1)[1]

    if action == "check":
        await callback.answer()
        await callback.message.reply_text(
            CHECK_PROMPT,
            reply_markup=ForceReply(
                selective=True,
                input_field_placeholder="Enter broker name or WikiFX ID",
            ),
        )
        return

    if action == "follows":
        if not update.effective_chat or update.effective_chat.type != "private":
            await callback.answer(
                "Open the bot’s private chat to manage followed brokers.",
                show_alert=True,
            )
            return
        await callback.answer()
        text, markup = follows_page(update.effective_user.id)
        await callback.edit_message_text(
            text,
            parse_mode=ParseMode.HTML,
            reply_markup=markup,
        )
        return

    if action == "home":
        await callback.answer()
        await callback.edit_message_text(
            welcome_text(),
            parse_mode=ParseMode.HTML,
            reply_markup=welcome_buttons(
                bool(update.effective_chat and update.effective_chat.type == "private"),
                context.bot.username or "",
            ),
        )


async def followed_broker_selected(update: Update, context: ContextTypes.DEFAULT_TYPE):
    callback = update.callback_query
    if not update.effective_chat or update.effective_chat.type != "private":
        await callback.answer("Open the bot’s private chat first.", show_alert=True)
        return

    broker_id = callback.data.split(":", 1)[1]
    records, _ = find_brokers(broker_id)
    if not records:
        await callback.answer("This broker is no longer active.", show_alert=True)
        return

    broker = records[0]
    await callback.answer()
    await callback.edit_message_text(
        broker_card(broker),
        parse_mode=ParseMode.HTML,
        reply_markup=broker_buttons(
            broker,
            user_id=update.effective_user.id,
            is_private=True,
            bot_username=context.bot.username or "",
        ),
        disable_web_page_preview=True,
    )


async def inline_query(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query_text = update.inline_query.query.strip()
    if not query_text:
        return
    records, match_type = find_brokers(query_text)
    results = []

    for broker in records[:5]:
        domain = next(iter(broker.get("domains_list", [])), "")
        results.append(
            InlineQueryResultArticle(
                id=str(uuid4()),
                title=str(broker["broker_name"]),
                description=(
                    f"{domain or 'Website not listed'} · "
                    f"{text_value(broker.get('country')) or 'Unknown region'} · "
                    f"ID {broker['broker_id']}"
                ),
                input_message_content=InputTextMessageContent(
                    broker_card(broker),
                    parse_mode=ParseMode.HTML,
                ),
                reply_markup=broker_buttons(
                    broker,
                    bot_username=context.bot.username or "",
                ),
            )
        )

    if not results:
        results.append(
            InlineQueryResultArticle(
                id=str(uuid4()),
                title=f"No exact match: {query_text}",
                description="Try the exact broker name, official website or 10-digit WikiFX ID",
                input_message_content=InputTextMessageContent(
                    "No match found. Try the exact broker name or 10-digit WikiFX ID."
                ),
            )
        )

    await update.inline_query.answer(results, cache_time=0, is_personal=True)


async def stats(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update):
        await update.message.reply_text("Admin access required.")
        return
    since = (datetime.now(timezone.utc) - timedelta(days=30)).isoformat()
    row = DB.execute(
        """
        SELECT
            SUM(CASE WHEN event_type='broker_query' THEN 1 ELSE 0 END) AS queries,
            COUNT(DISTINCT CASE WHEN event_type='broker_query' THEN telegram_user_id END) AS users,
            SUM(CASE WHEN event_type='match_success' THEN 1 ELSE 0 END) AS matches,
            SUM(CASE WHEN event_type='match_multiple' THEN 1 ELSE 0 END) AS multiple_matches,
            SUM(CASE WHEN event_type='match_none' THEN 1 ELSE 0 END) AS no_matches
        FROM events WHERE event_time >= ?
        """,
        (since,),
    ).fetchone()
    await update.message.reply_text(
        "India MVP — last 30 days\n\n"
        f"Queries: {row['queries'] or 0}\n"
        f"Unique users: {row['users'] or 0}\n"
        f"Single matches: {row['matches'] or 0}\n"
        f"Multiple matches: {row['multiple_matches'] or 0}\n"
        f"No matches: {row['no_matches'] or 0}\n\n"
        "Website visits, app installs and in-app deep visits are not included until analytics integration is complete."
    )


async def export_events(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update):
        await update.message.reply_text("Admin access required.")
        return
    cleanup_old_events()
    export_path = PROJECT_DIR / "telegram_events_last_30_days.csv"
    rows = DB.execute("SELECT * FROM events ORDER BY event_time DESC").fetchall()
    with export_path.open("w", newline="", encoding="utf-8-sig") as output:
        writer = csv.writer(output)
        writer.writerow(rows[0].keys() if rows else ["no_events"])
        for row in rows:
            writer.writerow(list(row))
    with export_path.open("rb") as document:
        await update.message.reply_document(document=document, filename=export_path.name)


async def push_news(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update):
        await update.message.reply_text("Admin access required.")
        return
    if len(context.args) < 2:
        await update.message.reply_text(
            "Admin usage: /pushnews <WikiFX ID> <news text or URL>"
        )
        return

    broker_id = context.args[0].strip()
    news_text = " ".join(context.args[1:]).strip()
    records, _ = find_brokers(broker_id)
    if len(records) != 1:
        await update.message.reply_text("Use one exact WikiFX ID for /pushnews.")
        return

    broker = records[0]
    followers = DB.execute(
        """
        SELECT telegram_user_id, private_chat_id
        FROM broker_follows
        WHERE broker_id = ? AND active = 1
        """,
        (str(broker["broker_id"]),),
    ).fetchall()
    if not followers:
        await update.message.reply_text("No active followers for this broker yet.")
        return

    sent = 0
    failed = 0
    sent_at = datetime.now(timezone.utc).isoformat()
    broker_name = text_value(broker.get("broker_name"))
    for follower in followers:
        try:
            await context.bot.send_message(
                chat_id=follower["private_chat_id"],
                text=(
                    f"📰 New update for {broker_name}\n\n"
                    f"{news_text}\n\n"
                    "Use /check to view the latest WikiFX broker snapshot."
                ),
                reply_markup=broker_buttons(
                    broker,
                    user_id=follower["telegram_user_id"],
                    is_private=True,
                    bot_username=context.bot.username or "",
                ),
                disable_web_page_preview=True,
            )
            DB.execute(
                """
                UPDATE broker_follows SET last_news_at = ?, updated_at = ?
                WHERE telegram_user_id = ? AND broker_id = ?
                """,
                (
                    sent_at,
                    sent_at,
                    follower["telegram_user_id"],
                    str(broker["broker_id"]),
                ),
            )
            sent += 1
        except TelegramError:
            failed += 1
    DB.commit()
    log_event(update, "admin_news_push", broker_id=broker["broker_id"], result_count=sent)
    await update.message.reply_text(
        f"News push complete for {broker_name}. Sent: {sent}; failed: {failed}."
    )


def recall_buttons(broker):
    broker_id = str(broker["broker_id"])
    broker_name = text_value(broker.get("broker_name"))[:22] or "broker"
    return InlineKeyboardMarkup(
        [
            [
                InlineKeyboardButton(
                    f"🔎 Check {broker_name} now",
                    callback_data=f"followview:{broker_id}",
                )
            ],
            [InlineKeyboardButton("⭐ My followed brokers", callback_data="nav:follows")],
            [InlineKeyboardButton("🔕 Stop these updates", callback_data=f"unfollow:{broker_id}")],
        ]
    )


def recall_text(broker_names):
    summary = ", ".join(broker_names)
    return (
        f"👋 <b>Still tracking {html.escape(summary)}?</b>\n\n"
        "A fresh WikiFX check can help you review the latest website, regulation and profile details. "
        "Tap below to refresh the broker card.\n\n"
        "<i>You are receiving this reminder because you followed broker updates.</i>"
    )


async def send_due_recalls(bot, inactivity_days: int, repeat_days: int):
    inactivity_days = max(1, min(inactivity_days, 30))
    repeat_days = max(1, min(repeat_days, 30))
    now = datetime.now(timezone.utc)
    activity_cutoff = now - timedelta(days=inactivity_days)
    recall_cutoff = now - timedelta(days=repeat_days)
    rows = DB.execute(
        """
        SELECT * FROM broker_follows
        WHERE active = 1
        ORDER BY telegram_user_id, updated_at DESC
        """
    ).fetchall()
    follows_by_user = defaultdict(list)
    for row in rows:
        follows_by_user[row["telegram_user_id"]].append(row)

    sent = 0
    skipped = 0
    failed = 0
    recalled_at = datetime.now(timezone.utc).isoformat()
    for user_id, follows in follows_by_user.items():
        last_activity_row = DB.execute(
            """
            SELECT MAX(event_time) AS last_activity FROM events
            WHERE telegram_user_id = ?
              AND event_type IN ('bot_start', 'broker_query', 'follow_enabled')
            """,
            (user_id,),
        ).fetchone()
        last_activity = last_activity_row["last_activity"] if last_activity_row else None
        last_recall = max(
            (row["last_recall_at"] for row in follows if row["last_recall_at"]),
            default=None,
        )
        if (last_activity and last_activity > activity_cutoff.isoformat()) or (
            last_recall and last_recall > recall_cutoff.isoformat()
        ):
            skipped += 1
            continue

        broker_ids = [row["broker_id"] for row in follows]
        broker_names = []
        first_broker = None
        for followed_broker_id in broker_ids[:3]:
            records, _ = find_brokers(followed_broker_id)
            if records:
                first_broker = first_broker or records[0]
                broker_names.append(text_value(records[0].get("broker_name")))
        if not first_broker:
            skipped += 1
            continue

        try:
            await bot.send_message(
                chat_id=follows[0]["private_chat_id"],
                text=recall_text(broker_names),
                parse_mode=ParseMode.HTML,
                reply_markup=recall_buttons(first_broker),
                disable_web_page_preview=True,
            )
            DB.execute(
                """
                UPDATE broker_follows SET last_recall_at = ?, updated_at = ?
                WHERE telegram_user_id = ? AND active = 1
                """,
                (recalled_at, recalled_at, user_id),
            )
            sent += 1
        except TelegramError:
            failed += 1
    DB.commit()
    return {"sent": sent, "skipped": skipped, "failed": failed}


async def recall_followers(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update):
        await update.message.reply_text("Admin access required.")
        return

    days = int(context.args[0]) if context.args and context.args[0].isdigit() else RECALL_AFTER_DAYS
    result = await send_due_recalls(context.bot, days, RECALL_REPEAT_DAYS)
    sent = result["sent"]
    log_event(update, "admin_recall", result_count=sent)
    await update.message.reply_text(
        "Follower recall complete. "
        f"Sent: {sent}; skipped: {result['skipped']}; failed: {result['failed']}."
    )


async def recall_preview(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update):
        return
    if not update.effective_chat or update.effective_chat.type != "private":
        await update.message.reply_text("Use /recallpreview in the bot’s private chat.")
        return
    if not context.args:
        await update.message.reply_text("Admin usage: /recallpreview <WikiFX ID>")
        return

    records, _ = find_brokers(context.args[0])
    if len(records) != 1:
        await update.message.reply_text("Use one exact WikiFX ID for the recall preview.")
        return
    broker = records[0]
    await update.message.reply_text(
        recall_text([text_value(broker.get("broker_name"))]),
        parse_mode=ParseMode.HTML,
        reply_markup=recall_buttons(broker),
    )


async def automatic_recall_loop(application: Application):
    await asyncio.sleep(min(60, max(5, RECALL_CHECK_SECONDS)))
    while True:
        try:
            result = await send_due_recalls(
                application.bot,
                RECALL_AFTER_DAYS,
                RECALL_REPEAT_DAYS,
            )
            if result["sent"] or result["failed"]:
                LOGGER.info(
                    "Automatic follower recall: sent=%s skipped=%s failed=%s",
                    result["sent"],
                    result["skipped"],
                    result["failed"],
                )
        except Exception:
            LOGGER.exception("Automatic follower recall check failed")
        await asyncio.sleep(max(300, RECALL_CHECK_SECONDS))


async def stop_background_tasks(application: Application):
    task = application.bot_data.get("automatic_recall_task")
    if task:
        task.cancel()
        try:
            await task
        except asyncio.CancelledError:
            pass


async def pin(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update) or not update.message.reply_to_message:
        await update.message.reply_text("Admin: reply to a message with /pin.")
        return
    await context.bot.pin_chat_message(update.effective_chat.id, update.message.reply_to_message.message_id)
    await update.message.reply_text("Message pinned.")


async def delete(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update) or not update.message.reply_to_message:
        await update.message.reply_text("Admin: reply to a message with /delete.")
        return
    await context.bot.delete_message(update.effective_chat.id, update.message.reply_to_message.message_id)


async def mute(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update) or not update.message.reply_to_message:
        await update.message.reply_text("Admin: reply to a user's message with /mute 60.")
        return
    minutes = int(context.args[0]) if context.args and context.args[0].isdigit() else 60
    until = datetime.now(timezone.utc) + timedelta(minutes=minutes)
    target = update.message.reply_to_message.from_user
    await context.bot.restrict_chat_member(
        update.effective_chat.id,
        target.id,
        permissions=ChatPermissions(can_send_messages=False),
        until_date=until,
    )
    await update.message.reply_text(f"Muted {target.full_name} for {minutes} minutes.")


async def unmute(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not is_admin(update) or not update.message.reply_to_message:
        await update.message.reply_text("Admin: reply to a user's message with /unmute.")
        return
    target = update.message.reply_to_message.from_user
    await context.bot.restrict_chat_member(
        update.effective_chat.id,
        target.id,
        permissions=ChatPermissions(
            can_send_messages=True,
            can_send_audios=True,
            can_send_documents=True,
            can_send_photos=True,
            can_send_videos=True,
            can_send_video_notes=True,
            can_send_voice_notes=True,
            can_send_polls=True,
            can_send_other_messages=True,
            can_add_web_page_previews=True,
            can_change_info=False,
            can_invite_users=True,
            can_pin_messages=False,
            can_manage_topics=False,
        ),
    )
    await update.message.reply_text(f"Unmuted {target.full_name}.")


def known_group_chat_ids():
    rows = DB.execute(
        """
        SELECT DISTINCT chat_id FROM events
        WHERE chat_id IS NOT NULL AND chat_type IN ('group', 'supergroup')
        """
    ).fetchall()
    return [int(row["chat_id"]) for row in rows]


async def configure_bot_commands(application: Application):
    try:
        await application.bot.set_my_commands(
            commands=PUBLIC_COMMANDS, scope=BotCommandScopeDefault()
        )
        await application.bot.set_my_commands(
            commands=PUBLIC_COMMANDS, scope=BotCommandScopeAllPrivateChats()
        )
        await application.bot.set_my_commands(
            commands=PUBLIC_COMMANDS, scope=BotCommandScopeAllGroupChats()
        )
    except TelegramError as error:
        LOGGER.warning("Could not configure the public command menu: %s", type(error).__name__)

    for admin_id in admin_ids():
        try:
            await application.bot.set_my_commands(
                commands=PUBLIC_COMMANDS,
                scope=BotCommandScopeChat(chat_id=admin_id),
            )
        except TelegramError as error:
            LOGGER.warning(
                "Could not configure private admin commands for user %s: %s",
                admin_id,
                type(error).__name__,
            )
        for chat_id in known_group_chat_ids():
            try:
                await application.bot.set_my_commands(
                    commands=PUBLIC_COMMANDS,
                    scope=BotCommandScopeChatMember(chat_id=chat_id, user_id=admin_id),
                )
            except TelegramError as error:
                LOGGER.warning(
                    "Could not configure admin commands for chat %s: %s",
                    chat_id,
                    type(error).__name__,
                )

    application.bot_data["automatic_recall_task"] = asyncio.create_task(
        automatic_recall_loop(application),
        name="wikifx-follower-recall",
    )


def main():
    token = os.getenv("TELEGRAM_BOT_TOKEN", "").strip()
    if not token:
        raise SystemExit("TELEGRAM_BOT_TOKEN is missing.")
    if not XLSX_PATH.exists():
        raise SystemExit(f"Broker workbook is missing: {XLSX_PATH}")

    cleanup_old_events()
    LOGGER.info("Using workbook: %s", XLSX_PATH)
    LOGGER.info("Loaded %s active WikiFX brokers", len(load_brokers()))
    LOGGER.info(
        "DeepSeek AI layer: %s",
        "enabled" if is_ai_enabled() else "disabled (missing DEEPSEEK_API_KEY)",
    )

    application = (
        Application.builder()
        .token(token)
        .post_init(configure_bot_commands)
        .post_shutdown(stop_background_tasks)
        .build()
    )
    application.add_handler(CommandHandler("start", start))
    application.add_handler(CommandHandler("help", start))
    application.add_handler(CommandHandler("privacy", privacy))
    application.add_handler(CommandHandler("check", check))
    application.add_handler(CommandHandler("stats", stats))
    application.add_handler(CommandHandler("export", export_events))
    application.add_handler(CommandHandler("pushnews", push_news))
    application.add_handler(CommandHandler("recall", recall_followers))
    application.add_handler(CommandHandler("recallpreview", recall_preview))
    application.add_handler(CommandHandler("pin", pin))
    application.add_handler(CommandHandler("delete", delete))
    application.add_handler(CommandHandler("mute", mute))
    application.add_handler(CommandHandler("unmute", unmute))
    application.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, text_message))
    application.add_handler(CallbackQueryHandler(candidate_selected, pattern=r"^broker:"))
    application.add_handler(
        CallbackQueryHandler(follow_selected, pattern=r"^(follow|unfollow):")
    )
    application.add_handler(CallbackQueryHandler(navigation_selected, pattern=r"^nav:"))
    application.add_handler(
        CallbackQueryHandler(followed_broker_selected, pattern=r"^followview:")
    )
    application.add_handler(InlineQueryHandler(inline_query))

    print("India MVP bot is running. Press Ctrl+C to stop.")
    application.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
