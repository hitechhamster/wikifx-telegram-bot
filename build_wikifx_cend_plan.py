from pathlib import Path
import math
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DOCX_PATH = ROOT / "WikiFX_C端交易商查询与召回运营方案_20260817.docx"
PNG_PATH = ROOT / "WikiFX_C端交易商查询闭环流程图_20260817.png"

BLUE = "1769AA"
DEEP_BLUE = "123B5D"
NAVY = "0B2545"
TEAL = "07847E"
LIGHT_BLUE = "EAF3FA"
LIGHT_TEAL = "E7F6F3"
LIGHT_ORANGE = "FFF1E6"
ORANGE = "E87524"
GOLD = "C58B16"
RED = "B93B37"
LIGHT_RED = "FDEDEC"
GRAY = "66717E"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "D9E0E7"
DARK = "1E2933"
WHITE = "FFFFFF"


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_run_font(run, name="Calibri", east_asia="Microsoft YaHei", size=None,
                 bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=MID_GRAY, size=6):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def add_bottom_border(paragraph, color=BLUE, size=12, space=8):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=GRAY)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, size=9, color=GRAY)


def add_hyperlink(paragraph, text, url, color=BLUE):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_fonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_paragraph(doc, text="", size=11, color=DARK, bold=False, italic=False,
                  align=WD_ALIGN_PARAGRAPH.LEFT, after=6, before=0, line=1.10,
                  keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if keep:
        pf.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(8 if level == 0 else 5)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=11, color=DARK)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=11, color=DARK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def add_callout(doc, title, body, fill=LIGHT_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_fill(cell, fill)
    set_cell_border(cell, color=accent, size=8)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=11, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, rows, widths, header_fill=LIGHT_GRAY, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_fill(cell, header_fill)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=font_size, bold=True, color=NAVY)
    for row_data in rows:
        row = table.add_row()
        set_cant_split(row)
        for i, text in enumerate(row_data):
            cell = row.cells[i]
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            r = p.add_run(str(text))
            set_run_font(r, size=font_size, color=DARK)
    return table


def add_table_citation(doc, text):
    p = add_paragraph(doc, text, size=8.5, color=GRAY, after=4, before=4, line=1.0)
    return p


def add_page_break(doc):
    doc.add_page_break()


def pil_font(size, bold=False):
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\arial.ttf",
    ]
    for p in candidates:
        if Path(p).exists():
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


def wrap_text(draw, text, font, max_width):
    lines = []
    for para in str(text).split("\n"):
        current = ""
        for ch in para:
            trial = current + ch
            if draw.textbbox((0, 0), trial, font=font)[2] <= max_width:
                current = trial
            else:
                if current:
                    lines.append(current)
                current = ch
        if current:
            lines.append(current)
    return lines


def draw_text_box(draw, box, title, body=None, fill="#FFFFFF", outline="#1769AA",
                  title_color="#123B5D", body_color="#334155", radius=26,
                  title_size=42, body_size=31, align="center", width=5):
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)
    tf = pil_font(title_size, bold=True)
    bf = pil_font(body_size, bold=False)
    title_lines = wrap_text(draw, title, tf, x2 - x1 - 56)
    body_lines = wrap_text(draw, body, bf, x2 - x1 - 56) if body else []
    th = sum((draw.textbbox((0, 0), t, font=tf)[3] - draw.textbbox((0, 0), t, font=tf)[1] + 8) for t in title_lines)
    bh = sum((draw.textbbox((0, 0), t, font=bf)[3] - draw.textbbox((0, 0), t, font=bf)[1] + 6) for t in body_lines)
    total = th + (18 if body_lines else 0) + bh
    y = y1 + (y2 - y1 - total) / 2
    for line in title_lines:
        bbox = draw.textbbox((0, 0), line, font=tf)
        w = bbox[2] - bbox[0]
        tx = x1 + 28 if align == "left" else x1 + (x2 - x1 - w) / 2
        draw.text((tx, y), line, font=tf, fill=title_color)
        y += bbox[3] - bbox[1] + 8
    if body_lines:
        y += 10
    for line in body_lines:
        bbox = draw.textbbox((0, 0), line, font=bf)
        w = bbox[2] - bbox[0]
        tx = x1 + 28 if align == "left" else x1 + (x2 - x1 - w) / 2
        draw.text((tx, y), line, font=bf, fill=body_color)
        y += bbox[3] - bbox[1] + 6


def draw_arrow(draw, start, end, color="#617487", width=9, head=24):
    x1, y1 = start
    x2, y2 = end
    draw.line((x1, y1, x2, y2), fill=color, width=width)
    angle = math.atan2(y2 - y1, x2 - x1)
    p1 = (x2 - head * math.cos(angle - math.pi / 6), y2 - head * math.sin(angle - math.pi / 6))
    p2 = (x2 - head * math.cos(angle + math.pi / 6), y2 - head * math.sin(angle + math.pi / 6))
    draw.polygon([(x2, y2), p1, p2], fill=color)


def draw_poly_arrow(draw, points, color="#617487", width=9, head=24):
    draw.line(points, fill=color, width=width, joint="curve")
    draw_arrow(draw, points[-2], points[-1], color=color, width=width, head=head)


def generate_flowchart():
    W, H = 2800, 3500
    img = Image.new("RGB", (W, H), "#F7FAFC")
    d = ImageDraw.Draw(img)

    d.rounded_rectangle((70, 55, W - 70, 330), radius=42, fill="#0B2545")
    d.text((140, 105), "WikiFX C端交易商查询与召回运营闭环", font=pil_font(78, True), fill="white")
    d.text((145, 220), "目标：每天新增100个可验证的APP/Web交易商深度访问（APP优先）", font=pil_font(39), fill="#CFE8F7")

    # Source row
    d.text((105, 390), "01 现有流量入口", font=pil_font(42, True), fill="#1769AA")
    sources = [
        (100, 475, 560, 690, "文章 / SEO / 新闻", "CTA：查一下这个交易商"),
        (610, 475, 1070, 690, "视频 / 社媒", "评论区、主页、私信入口"),
        (1120, 475, 1580, 690, "KOL现有私域", "不抢群，把工具装进群"),
        (1630, 475, 2090, 690, "站内现有用户", "交易商页/投诉页/风险页"),
        (2140, 475, 2700, 690, "投广", "只算获客来源，不算私域池"),
    ]
    for x1, y1, x2, y2, t, b in sources:
        draw_text_box(d, (x1, y1, x2, y2), t, b, fill="#FFFFFF", outline="#A8C5DB", title_size=35, body_size=27)

    draw_poly_arrow(d, [(1400, 720), (1400, 780), (760, 780), (760, 840)], color="#6A7B89")
    draw_poly_arrow(d, [(1400, 720), (1400, 780), (2040, 780), (2040, 840)], color="#6A7B89")

    # Channel lane
    d.text((105, 805), "02 聊天工具入口", font=pil_font(42, True), fill="#1769AA")
    draw_text_box(d, (170, 870, 1300, 1125), "WhatsApp 一对一查询助手（主路径）",
                  "用户点带来源码的链接，发送交易商名称；适合文章、现有用户与广泛C端触达",
                  fill="#EAF3FA", outline="#1769AA", title_size=43, body_size=31)
    draw_text_box(d, (1500, 870, 2630, 1125), "Telegram KOL群内机器人（盘外路径）",
                  "用户在KOL原有群中输入 @WikiFXCheckBot + 交易商；结果卡可被全群看到并继续分享",
                  fill="#E7F6F3", outline="#07847E", title_size=43, body_size=31)
    draw_arrow(d, (735, 1125), (735, 1225))
    draw_arrow(d, (2065, 1125), (2065, 1225))
    draw_text_box(d, (480, 1225, 2320, 1435), "用户输入交易商名称 / 官网链接",
                  "例：Exness、XM、ABC Markets；每个入口带 article_id / kol_id / market / campaign 来源参数",
                  fill="#FFFFFF", outline="#6A7B89", title_size=46, body_size=32)

    draw_arrow(d, (1400, 1435), (1400, 1520))
    draw_text_box(d, (480, 1520, 2320, 1740), "WikiFX查询层：匹配交易商数据库",
                  "自动版：API + Webhook；无API时先人工搜索并用固定模板回复，不等待产品改造",
                  fill="#FFF1E6", outline="#E87524", title_size=46, body_size=32)

    # Match branches
    draw_poly_arrow(d, [(1400, 1740), (1400, 1805), (690, 1805), (690, 1880)], color="#617487")
    draw_poly_arrow(d, [(1400, 1740), (1400, 1805), (1400, 1880)], color="#617487")
    draw_poly_arrow(d, [(1400, 1740), (1400, 1805), (2110, 1805), (2110, 1880)], color="#617487")
    draw_text_box(d, (145, 1880, 965, 2095), "唯一匹配",
                  "直接返回简版风险卡", fill="#E7F6F3", outline="#07847E", title_size=40, body_size=30)
    draw_text_box(d, (990, 1880, 1810, 2095), "多结果",
                  "先返回2—5个候选供用户点选", fill="#EAF3FA", outline="#1769AA", title_size=40, body_size=30)
    draw_text_box(d, (1835, 1880, 2655, 2095), "无结果",
                  "索要官网/截图 → 人工核验 → 结果后发", fill="#FDEDEC", outline="#B93B37", title_size=40, body_size=30)

    draw_poly_arrow(d, [(555, 2095), (555, 2170), (1135, 2170), (1135, 2235)], color="#07847E")
    draw_poly_arrow(d, [(1400, 2095), (1400, 2170), (1135, 2170), (1135, 2235)], color="#1769AA")
    draw_poly_arrow(d, [(2245, 2095), (2245, 2170), (1665, 2170), (1665, 2235)], color="#B93B37")
    draw_text_box(d, (470, 2235, 2330, 2515), "简版查询卡（聊天内给答案，但不替代完整详情）",
                  "交易商名称｜WikiFX评分｜注册地区｜监管状态｜核心风险｜更新时间\n按钮：完整牌照｜投诉记录｜风险证据｜关注这家交易商",
                  fill="#FFFFFF", outline="#123B5D", title_size=45, body_size=31)

    # Action split
    draw_poly_arrow(d, [(1400, 2515), (1400, 2585), (720, 2585), (720, 2655)], color="#1769AA")
    draw_poly_arrow(d, [(1400, 2515), (1400, 2585), (2080, 2585), (2080, 2655)], color="#07847E")
    draw_text_box(d, (120, 2655, 1320, 2880), "A. 立即查看完整证据",
                  "Smart Link判断：已安装 → APP精确交易商页；未安装 → Web精确交易商页，不先强制下载",
                  fill="#EAF3FA", outline="#1769AA", title_size=42, body_size=30)
    draw_text_box(d, (1480, 2655, 2680, 2880), "B. WATCH 关注交易商",
                  "记录用户×Broker ID与授权；进入高频内容召回或该交易商低频变更召回",
                  fill="#E7F6F3", outline="#07847E", title_size=42, body_size=30)

    draw_arrow(d, (720, 2880), (720, 2985), color="#1769AA")
    draw_poly_arrow(d, [(2080, 2880), (2080, 2940), (1680, 2940), (1680, 2985)], color="#07847E")
    draw_poly_arrow(d, [(2080, 2880), (2080, 2940), (2480, 2940), (2480, 2985)], color="#07847E")
    draw_text_box(d, (130, 2985, 1310, 3205), "即时转化",
                  "查询后24小时内进入APP/Web精确交易商深度页\n主指标：去重深访UV，APP与Web分开", fill="#FFFFFF", outline="#1769AA", title_size=42, body_size=30)
    draw_text_box(d, (1380, 2985, 1980, 3205), "高频召回",
                  "每日风险/监管提醒/热点交易商", fill="#FFFFFF", outline="#07847E", title_size=38, body_size=29)
    draw_text_box(d, (2180, 2985, 2780, 3205), "低频召回",
                  "所关注交易商的牌照/投诉/风险变化", fill="#FFFFFF", outline="#07847E", title_size=38, body_size=29)

    draw_poly_arrow(d, [(1680, 3205), (1680, 3260), (2080, 3260), (2080, 3310)], color="#07847E")
    draw_poly_arrow(d, [(2480, 3205), (2480, 3260), (2080, 3260), (2080, 3310)], color="#07847E")
    draw_text_box(d, (1450, 3310, 2710, 3440), "召回转化：提醒后进入精确交易商深度页",
                  "群成员数、消息送达和普通点击都不算北极星结果", fill="#0B2545", outline="#0B2545",
                  title_color="#FFFFFF", body_color="#CFE8F7", title_size=36, body_size=27)

    img.save(PNG_PATH, format="PNG", optimize=True, dpi=(240, 240))


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(DARK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DEEP_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_name in ("List Bullet", "List Number"):
        style = styles[list_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167
    if "List Bullet 2" in styles:
        style = styles["List Bullet 2"]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.left_indent = Inches(0.75)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(5)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = False

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("WikiFX｜C端交易商查询与召回运营方案")
    set_run_font(r, size=8.5, color=GRAY, bold=True)
    add_bottom_border(hp, color=MID_GRAY, size=4, space=5)

    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)


def set_standard_section_geometry(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def set_flowchart_section_geometry(section):
    # Named override: the full-page figure needs more vertical room.
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.20)
    section.footer_distance = Inches(0.20)


def build_document():
    doc = Document()
    configure_styles(doc)
    configure_document(doc)
    props = doc.core_properties
    props.title = "WikiFX C端交易商查询与召回运营方案"
    props.subject = "WhatsApp查询助手 + Telegram KOL群机器人；目标每天新增100个深度访问"
    props.author = "Carmy"
    props.keywords = "WikiFX, WhatsApp, Telegram, C端, 交易商查询, 召回, 深度访问"

    # Cover / memo masthead
    add_paragraph(doc, "WIKIFX · C端增长运营", size=10, color=BLUE, bold=True, after=5, keep=True)
    p = add_paragraph(doc, "交易商查询与召回运营方案", size=25, color=NAVY, bold=True, after=4, keep=True)
    add_paragraph(doc, "WhatsApp一对一查询助手 + Telegram KOL群内机器人", size=14, color=DEEP_BLUE, bold=False, after=16, keep=True)

    meta_rows = [
        ("目标", "每天新增100个可验证的APP/Web交易商深度访问，APP优先"),
        ("范围", "只做C端；印度、巴基斯坦、印度尼西亚；不以产品改版为前提"),
        ("负责人", "Carmy｜内容与自然增长运营"),
        ("版本", "内部汇报稿｜2026年8月17日"),
    ]
    table = doc.add_table(rows=0, cols=2)
    set_table_geometry(table, [1800, 7560])
    for label, value in meta_rows:
        row = table.add_row()
        for cell in row.cells:
            set_cell_border(cell, color="FFFFFF", size=0)
        set_cell_fill(row.cells[0], LIGHT_BLUE)
        set_cell_fill(row.cells[1], "FFFFFF")
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(label)
        set_run_font(r1, size=10, bold=True, color=BLUE)
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(value)
        set_run_font(r2, size=10, color=DARK)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    add_bottom_border(rule, color=BLUE, size=14, space=7)

    add_callout(
        doc,
        "一句话方案",
        "把WikiFX交易商数据库变成聊天场景里的“查询入口”：用户先在WhatsApp或Telegram输入交易商名称，收到简版风险结果；当他要看完整牌照、投诉和证据时，再进入APP或Web的精确交易商页。查询解决即时转化，WATCH关注解决高频/低频召回。",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )
    add_heading(doc, "领导这次真正要的是什么", 1)
    add_bullet(doc, "不是再建一个“内容群”，而是找到一个可持续把用户带回WikiFX深度页的运营机制。")
    add_bullet(doc, "不是用下载弹窗硬推未安装用户，而是先满足查询，再让高意向用户自然进入APP。")
    add_bullet(doc, "不是看群人数或消息点击，而是每天能新增多少可归因、可验证的交易商深访UV。")
    add_bullet(doc, "盘外招不是把KOL用户抢进新群，而是让WikiFX机器人直接进入KOL现有Telegram群。")

    add_heading(doc, "建议结论", 1)
    add_callout(
        doc,
        "推荐采用“1主1辅”",
        "主路径：WhatsApp一对一交易商查询助手，覆盖文章、现有用户、视频、社媒与投广入口。辅路径：Telegram群内Inline Bot，嵌入KOL现有社群，查询结果在群内自然曝光和分享。第一阶段先人工闭环，不等接口、不改APP页面。",
        fill=LIGHT_TEAL,
        accent=TEAL,
    )

    add_page_break(doc)
    add_heading(doc, "1. 项目目标与边界", 1)
    add_heading(doc, "1.1 北极星指标", 2)
    add_table(
        doc,
        ["指标", "定义", "为什么"],
        [
            ("每日新增深访UV", "相对基线新增的APP/Web精确交易商深度页去重访问；目标+100/天", "领导最终要看到的结果"),
            ("APP深访UV", "进入APP具体交易商详情、牌照、投诉或风险证据页的去重用户", "优先指标"),
            ("Web深访UV", "未安装用户进入Web具体交易商详情页的去重用户", "承接未安装用户，不浪费意图"),
            ("无效结果", "群成员、消息送达、普通落地页点击、打开下载页但未进入深度页", "只能做过程指标"),
        ],
        [1900, 4800, 2660],
    )
    add_heading(doc, "1.2 四条边界", 2)
    add_bullet(doc, "只考虑C端用户，B端线索与低频B端需求不纳入本方案。")
    add_bullet(doc, "第一阶段不依赖产品页面改造；优先复用现有Web交易商页、APP Deep Link和运营工具。")
    add_bullet(doc, "私域不是一个新群，而是可识别用户、可记录兴趣、可获得授权、可重复触达的关系池。")
    add_bullet(doc, "投广属于获客来源，可以进入同一闭环，但广告受众本身不等于私域池。")

    add_heading(doc, "1.3 核心判断", 2)
    add_callout(doc, "本方案最重要的产品化思维", "聊天工具负责“问问题、拿简版答案、订阅变化”；WikiFX APP/Web负责“看完整证据、完成深度访问”。聊天端不应该把全部内容讲完，否则反而会吃掉APP深访。", fill=LIGHT_ORANGE, accent=ORANGE)

    add_page_break(doc)
    add_heading(doc, "2. 市场与竞品调研：同行做到了哪一步", 1)
    add_paragraph(doc, "公开可见的直接竞品大多仍以网页交易商查询、评分、对比和Telegram内容群为主；成熟的“聊天内查询—简版结果—站内深访”能力，更多出现在银行客服、反诈工具和行情工具中。", after=10)
    add_table(
        doc,
        ["产品/行业", "公开做法", "可复制点", "对WikiFX的判断"],
        [
            ("BrokerChooser", "网站Verify broker、经纪商匹配与对比", "查询意图强，入口表述直接", "可复制“验证交易商”CTA，但仍是Web内闭环"),
            ("Traders Union", "网站Check Broker / Best Broker for Me", "搜索+推荐+比较", "说明用户愿意主动输入交易商"),
            ("BrokersView", "Telegram群提供资讯、评论、排名、假平台和监管提醒", "内容触达和群体讨论", "公开页面未见成熟群内自动查询机器人"),
            ("Forex Peace Army", "网站搜索、评论、维权帖子与论坛", "投诉和证据是深访动机", "公开页面未见聊天机器人查询路径"),
            ("ICICI / HDFC银行", "WhatsApp中输入关键词或自然语言，返回快捷服务；支持订阅/退订", "用户不需要学习复杂命令", "最适合复制到一对一交易商查询与召回"),
            ("ScamAdviser", "Telegram机器人可加入群，自动扫描URL并返回Trust Score", "工具进入现有群、结果全群可见", "最接近WikiFX“群内输入/链接→简版风险→完整详情”"),
            ("TradingView", "Telegram Mini App中看图表、价格并分享快照", "群内工具化和可分享", "适合未来增强，MVP不做，避免过重"),
        ],
        [1700, 2750, 2300, 2610],
        font_size=8.7,
    )
    add_table_citation(doc, "说明：以上为公开页面与公开Telegram入口调研；“未见机器人”只代表公开可见信息，不排除私有试验。")

    add_heading(doc, "2.1 调研结论", 2)
    add_number(doc, "直接竞品已经证明“交易商查询、监管验证、投诉证据”是强需求，但大部分仍停留在网站。")
    add_number(doc, "WhatsApp银行客服证明：C端用户能理解“输入关键词/名称→即时结果”的低门槛交互。")
    add_number(doc, "ScamAdviser证明：机器人进入他人的Telegram群、自动返回风险结果，是可落地的群内工具模式。")
    add_number(doc, "WikiFX的机会不是再复制内容群，而是把独有的交易商数据做成可嵌入各类流量和社群的查询工具。")

    add_page_break(doc)
    add_heading(doc, "3. 核心方案：查询解决即时转化，关注解决召回转化", 1)
    add_heading(doc, "3.1 方案架构", 2)
    add_table(
        doc,
        ["模块", "用户动作", "运营动作", "产生的结果"],
        [
            ("流量入口", "看到文章、视频、KOL内容、社媒或广告", "为每个入口配置独立来源码", "知道用户从哪里来"),
            ("聊天查询", "输入交易商名称/网址", "人工或机器人匹配WikiFX数据库", "形成有效查询"),
            ("简版结果", "阅读评分、监管与核心风险", "只给决策摘要，保留完整证据", "建立信任和深访动机"),
            ("立即深访", "点完整牌照/投诉/风险证据", "Smart Link分流到APP精确页或Web精确页", "即时转化"),
            ("WATCH关注", "授权关注某个交易商", "保存User ID × Broker ID × 来源", "形成可运营关系池"),
            ("召回", "收到热点或具体交易商变化提醒", "高频/低频两类内容，均跳精确深度页", "召回转化"),
        ],
        [1600, 2650, 3100, 2010],
        font_size=9.0,
    )
    add_heading(doc, "3.2 用户在聊天内能拿到什么", 2)
    add_callout(
        doc,
        "简版查询卡示例",
        "查询结果：ABC Markets\nWikiFX评分：X.X/10｜注册地区：XX\n监管状态：受监管 / 离岸监管 / 暂未验证\n核心提醒：存在X项风险提示｜数据更新时间：2026-XX-XX\n操作：查看完整牌照｜查看投诉记录｜查看风险证据｜关注这家交易商",
        fill="FFFFFF",
        accent=DEEP_BLUE,
    )
    add_paragraph(doc, "原则：简版卡必须“有用但不讲完”。用户能在聊天内完成初步判断，完整牌照文件、监管证据、投诉详情与历史变化留在WikiFX精确页面。", size=10.5, color=GRAY)

    chart_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    set_flowchart_section_geometry(chart_section)
    add_heading(doc, "4. 总体流程图", 1)
    add_paragraph(doc, "下图把两类入口、三种匹配结果、即时转化与高低频召回串成一个可归因闭环。", after=8)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.keep_with_next = True
    run = p_img.add_run()
    run.add_picture(str(PNG_PATH), width=Inches(6.5))
    caption = add_paragraph(doc, "图1｜WikiFX C端交易商查询与召回运营闭环（高清PNG已单独输出）", size=9, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=4, keep=True)
    caption.paragraph_format.keep_with_next = True
    add_callout(doc, "图中最关键的分流", "已安装用户直接打开APP精确交易商页；未安装用户先落到Web精确交易商页，而不是先进入下载页。只有在用户已查看风险/牌照并显示明确意图后，才二次邀请安装APP。", fill=LIGHT_ORANGE, accent=ORANGE)

    standard_section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    set_standard_section_geometry(standard_section)
    add_heading(doc, "5. WhatsApp主路径：一对一交易商查询助手", 1)
    add_heading(doc, "5.1 用户实操路径", 2)
    steps = [
        ("1", "进入", "用户从文章CTA、WikiFX页面、视频简介、KOL内容或广告点击带来源码的WhatsApp链接。"),
        ("2", "提问", "输入交易商名称、官网链接，或发送“CHECK + 交易商”。链接可预填示例话术。"),
        ("3", "匹配", "人工MVP由运营在WikiFX搜索；自动版由Webhook调用交易商搜索接口。"),
        ("4", "返回", "唯一匹配给简版卡；多结果给候选列表；无结果索要官网/截图并转人工。"),
        ("5", "深访", "用户点牌照/投诉/风险按钮，进入APP或Web的精确交易商页面。"),
        ("6", "关注", "用户回复WATCH并完成授权，记录User ID × Broker ID，后续按变化召回。"),
    ]
    add_table(doc, ["步骤", "动作", "具体表现"], steps, [900, 1300, 7160], font_size=9.5)

    add_heading(doc, "5.2 三种匹配规则", 2)
    add_table(
        doc,
        ["情况", "回复逻辑", "避免的问题"],
        [
            ("唯一匹配", "立即返回简版卡与完整详情按钮", "减少等待和跳失"),
            ("同名/多结果", "返回2—5个候选：名称、地区、官网域名，让用户点选", "防止查错主体"),
            ("无结果", "索要官网URL/截图，建立待核验任务；核验后主动回发", "把“查不到”变成低频召回机会"),
        ],
        [1700, 4700, 2960],
    )

    add_heading(doc, "5.3 现有能力与所需能力", 2)
    add_table(
        doc,
        ["能力", "人工MVP", "自动化版本"],
        [
            ("接待", "WhatsApp Business App欢迎语、快捷回复、标签", "WhatsApp Business Platform + Webhook"),
            ("查询", "运营人工在WikiFX搜索", "交易商搜索API/数据接口"),
            ("结果", "固定文本模板+现有详情链接", "动态卡片、候选列表、Reply Button"),
            ("关注", "飞书表记录用户、Broker ID、来源、授权", "数据库订阅关系 + 自动通知"),
            ("归因", "独立短链/UTM + 现有GA4/App事件", "完整事件链与用户级归因"),
        ],
        [2100, 3630, 3630],
    )

    add_page_break(doc)
    add_heading(doc, "6. Telegram辅路径：把工具放进KOL现有群", 1)
    add_heading(doc, "6.1 为什么这是“盘外招”", 2)
    add_callout(doc, "不是让KOL把用户交出来", "KOL不需要解散自己的群，也不需要把成员搬到WikiFX新群。WikiFX提供一个群内可用的交易商查询工具，KOL获得免费工具和内容素材，WikiFX获得带kol_id的查询和深访。", fill=LIGHT_TEAL, accent=TEAL)

    add_heading(doc, "6.2 群内具体操作", 2)
    add_number(doc, "KOL把 @WikiFXCheckBot 加入原有Telegram群，并固定一条使用说明。")
    add_number(doc, "用户在任何对话中输入“@WikiFXCheckBot Exness”或转发交易商官网链接。")
    add_number(doc, "机器人返回可选择的交易商结果；用户点选后，群内显示简版风险卡。")
    add_number(doc, "卡片提供完整牌照、投诉、风险证据和WATCH按钮；深访链接携带kol_id、group_id、broker_id。")
    add_number(doc, "其他群成员看到结果后可继续查询或分享，形成“一个查询带来多次曝光”。")

    add_heading(doc, "6.3 功能优先级", 2)
    add_table(
        doc,
        ["功能", "MVP优先级", "原因"],
        [
            ("Bot私聊查询", "P0", "最简单、最容易验证查询与深访"),
            ("Inline Bot群内查询", "P0/P1", "真正发挥KOL现有群的传播价值"),
            ("结果卡分享", "P1", "把单人查询变成群体曝光"),
            ("WATCH与变化通知", "P1", "建立低频召回能力"),
            ("Telegram Mini App", "P2", "能力重、开发成本高，且可能把用户留在Telegram而不回WikiFX"),
        ],
        [2800, 1600, 4960],
    )

    add_heading(doc, "6.4 KOL合作交换价值", 2)
    add_bullet(doc, "给KOL：群内免费交易商查询、风险结果卡、热点交易商素材、专属数据周报。")
    add_bullet(doc, "给WikiFX：真实查询词、精确深访、可衡量的KOL归因，以及新用户持续关注关系。")
    add_bullet(doc, "给用户：不用离开聊天先获得简版判断，需要完整证据时再进入WikiFX。")

    add_page_break(doc)
    add_heading(doc, "7. 未安装用户：不粗暴推下载，先承接意图", 1)
    add_heading(doc, "7.1 推荐分流", 2)
    add_table(
        doc,
        ["用户状态", "第一次点击", "第二次引导", "原则"],
        [
            ("已安装APP", "直接打开APP对应交易商详情/牌照/投诉页", "引导WATCH或查看更多证据", "减少中间页"),
            ("未安装APP", "打开Web对应交易商详情页", "在浏览完整证据、关注或二次回访时再提示APP", "先兑现价值，不拦截"),
            ("无法识别", "默认Web精确页", "提供“在APP中打开”入口", "保证任何用户都能看结果"),
        ],
        [1800, 3100, 3000, 1460],
    )

    add_heading(doc, "7.2 什么时候才适合邀请安装", 2)
    add_bullet(doc, "用户连续查看多个交易商或多个证据页。")
    add_bullet(doc, "用户点击WATCH，希望接收所关注交易商的变化提醒。")
    add_bullet(doc, "用户从召回消息第二次进入WikiFX，已证明有持续需求。")
    add_bullet(doc, "APP确实能提供Web没有的体验，例如推送、收藏、历史对比或更完整材料。")

    add_callout(doc, "关键文案", "不要说“下载APP才能查看”。建议说：“完整监管文件和历史变化可在WikiFX查看；已安装将直接打开APP，未安装也可先在网页查看。”", fill=LIGHT_ORANGE, accent=ORANGE)

    add_heading(doc, "7.3 技术实现顺序", 2)
    add_number(doc, "先确认现有Web是否有稳定的交易商详情URL。")
    add_number(doc, "确认现有APP是否支持Universal Link / Android App Link / AppsFlyer OneLink等精确Deep Link。")
    add_number(doc, "已具备则直接复用；不具备时，MVP全部先落Web精确页，APP引导放在页面内部。")
    add_number(doc, "Deferred Deep Link放到后续阶段：安装完成后仍能回到原交易商页，而不是APP首页。")

    add_page_break(doc)
    add_heading(doc, "8. 即时转化与召回转化", 1)
    add_heading(doc, "8.1 即时转化", 2)
    add_table(
        doc,
        ["定义", "触发内容", "落地页", "核心指标"],
        [
            ("用户完成有效交易商查询后24小时内进入精确深度页", "完整牌照、投诉记录、风险证据", "APP/Web交易商详情子页", "即时深访UV ÷ 有效查询UV"),
        ],
        [2600, 2400, 2500, 1860],
    )

    add_heading(doc, "8.2 召回转化：高频与低频必须拆开", 2)
    add_table(
        doc,
        ["类型", "面向谁", "内容例子", "发送规则", "衡量"],
        [
            ("高频召回", "授权接收市场风险/热点的用户", "今日风险交易商、监管警示、投诉上涨、热门交易商查询榜", "控制频次；每条只带一个明确交易商或榜单深访动作", "高频召回深访UV ÷ 高频送达UV"),
            ("低频召回", "WATCH某个交易商的用户", "该交易商牌照变化、评分变化、投诉增长、仿冒/克隆、实勘或重大舆情", "只有真实变化才发；附变化时间和证据", "低频召回深访UV ÷ 低频送达UV"),
        ],
        [1500, 1750, 3000, 1750, 1360],
        font_size=8.8,
    )

    add_heading(doc, "8.3 用户授权与退出", 2)
    add_bullet(doc, "WATCH前明确告知会接收什么、频率大致如何，并保留取消关注/UNSUB入口。")
    add_bullet(doc, "高频内容与低频交易商变化分开订阅，避免把所有人都塞进同一种触达。")
    add_bullet(doc, "手机号、WhatsApp ID或Telegram User ID只保留完成查询、召回和归因所需的最小字段。")
    add_bullet(doc, "任何触达都不能承诺收益、诱导入金或提供个性化投资建议。")

    add_page_break(doc)
    add_heading(doc, "9. 归因与数据闭环", 1)
    add_heading(doc, "9.1 最小事件链", 2)
    add_table(
        doc,
        ["阶段", "事件", "必须字段"],
        [
            ("入口", "entry_exposure / chat_start", "market, source_type, article_id/kol_id, campaign_id"),
            ("查询", "broker_query / match_success", "channel, user_id_hash, broker_keyword, broker_id, match_type"),
            ("结果", "summary_sent / detail_click", "broker_id, button_type, timestamp"),
            ("即时深访", "app_deep_view / web_deep_view", "broker_id, page_type, source_id, unique_user"),
            ("关注", "watch_opt_in / watch_cancel", "broker_id, consent_type, consent_time"),
            ("召回", "recall_sent / recall_deep_view", "recall_type, change_id/content_id, broker_id"),
        ],
        [1500, 3050, 4810],
        font_size=9.0,
    )

    add_heading(doc, "9.2 归因方法", 2)
    add_bullet(doc, "WhatsApp：每篇文章、每个KOL、每个市场使用独立预填文本或短链参数。")
    add_bullet(doc, "Telegram：Bot Deep Link使用start参数；群内链接增加kol_id、group_id和broker_id。")
    add_bullet(doc, "APP/Web：最终精确页面继续携带来源参数，并在GA4/App事件中记录broker_id和page_type。")
    add_bullet(doc, "同一用户同一日多次进入同一交易商深度页，北极星统计按去重UV；点击次数另列。")

    add_heading(doc, "9.3 汇报看板只看四块", 2)
    add_table(
        doc,
        ["结果块", "核心数", "拆分维度"],
        [
            ("新增深访", "相对基线新增APP UV + Web UV", "市场、渠道、交易商、来源"),
            ("即时效率", "有效查询→精确深访转化率", "WhatsApp / Telegram，APP / Web"),
            ("召回效率", "高频召回率、低频召回率", "触达内容、Broker ID、发送批次"),
            ("增长质量", "APP占比、重复查询率、WATCH率、取消率", "新/老用户、来源、市场"),
        ],
        [2100, 3800, 3460],
    )
    add_callout(doc, "验收底线", "WhatsApp加入、Telegram群曝光、消息送达和短链点击都不是最终结果。只有后端或分析工具确认用户进入了APP/Web精确交易商深度页，才计入北极星。", fill=LIGHT_RED, accent=RED)

    add_page_break(doc)
    add_heading(doc, "10. 最小可行试验：先人工跑通，不等开发", 1)
    add_heading(doc, "10.1 7—14天小范围闭环", 2)
    add_table(
        doc,
        ["阶段", "需要准备", "运营执行", "验证问题"],
        [
            ("准备", "1个官方WhatsApp Business号、固定回复模板、飞书记录表、可归因精确链接", "选20—30个高频交易商，整理简版答案和三类深访链接", "能否在5分钟内稳定回复"),
            ("导流", "2—3篇文章CTA、1个KOL、1个站内入口", "统一话术“发送交易商名称，免费查监管与风险”", "用户是否愿意主动输入"),
            ("即时", "唯一/多结果/无结果三套SOP", "人工查询、发简版卡、发精确页链接", "简版结果能否带来深访"),
            ("关注", "WATCH字段与授权话术", "记录User × Broker ID × 来源", "用户是否愿意关注变化"),
            ("召回", "1次高频内容 + 1次真实低频变化", "只发明确证据和精确深访链接", "哪类召回更有效"),
            ("复盘", "GA4/App事件与人工表", "核对来源、查询、点击、深访、WATCH、取消", "+100目标需要多少入口量"),
        ],
        [1200, 2800, 3300, 2060],
        font_size=8.7,
    )

    add_heading(doc, "10.2 首轮不追求什么", 2)
    add_bullet(doc, "不追求一次覆盖三个市场所有交易商；先覆盖高频和高风险查询。")
    add_bullet(doc, "不追求完整机器人；先证明用户会问、会点、会关注、会回来。")
    add_bullet(doc, "不追求群规模；先验证每100个有效查询能产生多少即时和召回深访。")
    add_bullet(doc, "不追求Mini App或复杂WhatsApp Flows；这些只在闭环成立后再开发。")

    add_heading(doc, "10.3 试验通过标准", 2)
    add_table(
        doc,
        ["问题", "最低应回答的证据"],
        [
            ("用户会不会输入交易商？", "有效查询UV、查询完成率、无结果率"),
            ("简版结果会不会带来深访？", "detail_click与APP/Web deep_view的转化"),
            ("用户会不会接受WATCH？", "WATCH率、取消率、关注Broker分布"),
            ("召回能不能带回深访？", "高频/低频分别统计recall_deep_view"),
            ("能否做到+100/天？", "按真实转化率倒推所需入口曝光和有效查询量"),
        ],
        [3400, 5960],
    )

    add_page_break(doc)
    add_heading(doc, "11. 分阶段落地与依赖", 1)
    add_table(
        doc,
        ["阶段", "周期", "范围", "是否需要产品/开发", "产出"],
        [
            ("阶段A｜人工MVP", "7—14天", "WhatsApp Business + 固定模板 + 飞书表 + 精确链接", "原则上不需要产品改版；需数据/链接协助", "真实转化率与问题清单"),
            ("阶段B｜半自动", "2—4周", "关键词解析、候选列表、结果卡、自动来源记录", "需要后端/接口与WhatsApp Platform", "降低人力、扩大查询量"),
            ("阶段C｜Telegram群内", "并行小试", "Bot私聊 + Inline Bot + KOL群", "需要Bot与查询接口", "验证群内传播与KOL归因"),
            ("阶段D｜自动召回", "闭环后", "Broker变化监测、订阅规则、自动通知", "需要变更数据与消息系统", "规模化低频召回"),
            ("阶段E｜深链优化", "闭环后", "Universal/App Link、Deferred Deep Link", "需要APP/归因工具配合", "提高APP深访占比"),
        ],
        [1750, 1150, 2860, 2450, 1150],
        font_size=8.6,
    )

    add_heading(doc, "11.1 启动前必须确认的五件事", 2)
    add_number(doc, "WikiFX现有交易商搜索/详情数据是否有可用API；没有则明确人工MVP响应时效。")
    add_number(doc, "Web交易商详情、牌照、投诉、风险证据是否有稳定且可直接访问的URL。")
    add_number(doc, "APP是否已有可打开精确交易商页的Deep Link，以及AppsFlyer/Adjust等归因能力。")
    add_number(doc, "WhatsApp Business官方账号、模板消息、24小时会话与用户授权的合规边界。")
    add_number(doc, "GA4/App事件能否记录broker_id、page_type、source_id并按去重用户出数。")

    add_heading(doc, "11.2 角色建议", 2)
    add_table(
        doc,
        ["角色", "首轮责任", "验收"],
        [
            ("Carmy/运营", "入口、话术、回复SOP、WATCH记录、召回内容、日报", "查询/深访/召回数据可对齐"),
            ("内容/SEO/KOL", "在现有内容与社群加入统一CTA和来源码", "每个入口可单独归因"),
            ("数据", "确认基线、事件、去重口径和日报字段", "能回答每天新增多少深访"),
            ("产品/开发", "仅在MVP验证后提供API、Bot、Deep Link优化", "自动化不改变验证逻辑"),
            ("法务/合规", "审核授权、退订、隐私和风险文案", "无收益承诺、无诱导入金"),
        ],
        [1800, 4800, 2760],
    )

    add_page_break(doc)
    add_heading(doc, "12. 风险、对策与最终建议", 1)
    add_table(
        doc,
        ["风险", "表现", "对策"],
        [
            ("查不到/查错", "小交易商、同名主体、网址变化", "多结果确认；无结果转人工；以Broker ID绑定"),
            ("回复慢", "人工MVP高峰积压", "先覆盖高频交易商；设5—15分钟SLA；沉淀标准答案"),
            ("聊天端讲得太全", "用户不再进入WikiFX", "只给结论摘要；完整证据必须落精确页"),
            ("未安装流失", "先跳下载页导致退出", "默认Web精确页；二次高意向再推APP"),
            ("召回骚扰", "退订、拉黑、投诉", "明确授权；高/低频分订阅；控制频率；随时UNSUB"),
            ("数据虚高", "把点击、群人数当深访", "以APP/Web deep_view去重事件验收"),
            ("KOL不愿导流", "担心流失用户", "工具留在原群；给KOL专属结果卡和周报"),
        ],
        [1900, 3000, 4460],
        font_size=9.0,
    )

    add_heading(doc, "12.1 最终建议", 2)
    add_callout(
        doc,
        "汇报时可直接这样说",
        "直接竞品大多还在做网页交易商查询和Telegram内容群。我们的突破口不是再建一个群，而是把WikiFX数据库做成WhatsApp查询助手和Telegram群内机器人：用户在聊天里输入交易商，先拿简版风险，再进入APP看完整证据；关注交易商后，再通过高频热点和低频真实变化把人带回来。首轮用人工方式7—14天即可验证，不等产品改版。",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )
    add_heading(doc, "12.2 汇报决策请求", 2)
    add_number(doc, "同意先做WhatsApp人工MVP，并指定1个官方号码。")
    add_number(doc, "同意选1个KOL做Telegram群内查询试点，不要求迁群。")
    add_number(doc, "数据/产品协助确认精确URL、Deep Link、GA4/App事件和API现状。")
    add_number(doc, "以“每天新增精确交易商深访UV”作为唯一结果验收，高频/低频召回分别出数。")

    add_page_break(doc)
    add_heading(doc, "附录A｜公开调研依据", 1)
    sources = [
        ("BrokerChooser｜Verify broker与经纪商匹配", "https://brokerchooser.com/"),
        ("Traders Union｜Check Broker", "https://tradersunion.com/"),
        ("BrokersView｜Telegram公开群", "https://t.me/Brokersview"),
        ("Forex Peace Army｜Reviews & Forums", "https://www.forexpeacearmy.com/"),
        ("ICICI Bank｜WhatsApp Banking", "https://www.icici.bank.in/about-us/article/news-icici-bank-launches-banking-services-on-whatsapp-20203003110358746"),
        ("HDFC Bank｜ChatBanking", "https://www.hdfc.bank.in/ways-to-bank/chat-banking"),
        ("HDFC｜SUB/UNSUB操作说明", "https://v.hdfc.bank.in/htdocs/common/whatsapp-banking/index.html"),
        ("ScamAdviser｜Telegram Bot", "https://www.scamadviser.com/articles/scamadviser-launches-at-scam-adviser_bot-to-fight-scams-on-telegram"),
        ("TradingView｜Telegram Mini App", "https://www.tradingview.com/blog/en/tradingview-mini-app-on-telegram-51366/"),
        ("Telegram｜Bot Features", "https://core.telegram.org/bots/features"),
        ("Telegram｜Mini Apps", "https://core.telegram.org/bots/webapps"),
        ("Meta｜WhatsApp Business Platform", "https://www.postman.com/meta/whatsapp-business-platform/overview"),
        ("Apple｜Universal Links", "https://developer.apple.com/documentation/xcode/allowing-apps-and-websites-to-link-to-your-content/"),
        ("Android｜App Links", "https://developer.android.com/training/app-links/about"),
        ("AppsFlyer｜OneLink", "https://dev.appsflyer.com/hc/docs/dl_getting_started"),
        ("WikiFX｜Google Play", "https://play.google.com/store/apps/details?id=com.foreigncurrency.internationalfxeye"),
    ]
    for i, (label, url) in enumerate(sources, 1):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(label + "：")
        set_run_font(r, size=9.5, color=DARK)
        add_hyperlink(p, url, url)

    add_heading(doc, "附录B｜首轮固定回复模板", 1)
    add_callout(
        doc,
        "欢迎语",
        "你好，这里是WikiFX交易商查询助手。请发送交易商名称或官网链接，我会先给你一份简版监管与风险结果；完整牌照、投诉和风险证据可继续打开WikiFX查看。回复 STOP 可停止接收消息。",
        fill=LIGHT_GRAY,
        accent=GRAY,
    )
    add_callout(
        doc,
        "查询结果模板",
        "查询结果：{Broker Name}\nWikiFX评分：{Score}/10｜注册地区：{Region}\n监管状态：{Status}\n核心提醒：{Risk Summary}\n数据更新时间：{Updated At}\n请选择：完整牌照｜投诉记录｜风险证据｜WATCH关注",
        fill=LIGHT_BLUE,
        accent=BLUE,
    )
    add_callout(
        doc,
        "无结果模板",
        "暂未找到完全匹配的交易商。请发送其官网地址或截图，我们会继续核验。核验完成后可直接回发结果；若你同意，我们也可以在结果有更新时通知你。",
        fill=LIGHT_RED,
        accent=RED,
    )

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    generate_flowchart()
    build_document()
    print(DOCX_PATH)
    print(PNG_PATH)
